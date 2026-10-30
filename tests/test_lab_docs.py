"""Laboratory Docs: LAB_STAFF-only access and Word/CSV report generation (no Snowflake needed)."""

import io

import pandas as pd
import pytest

from mediverse import lab_docs as ld
from mediverse.authz import AccessDenied, Principal, Role

HDR = {"PATIENT_ID": "P000001", "NAME": "Test Patient", "DATE_OF_BIRTH": "1980-01-01", "SEX": "F", "BLOOD_TYPE": "O+"}
RES = pd.DataFrame([
    {"LAB_ORDER_ID": "LO1", "ORDERED_TS": pd.Timestamp("2025-01-02"), "PRIORITY": "ROUTINE", "TEST_CODE": "HGB",
     "TEST_NAME": "Hemoglobin", "VALUE": 9.1, "UNIT": "g/dL", "REF_LOW": 12.0, "REF_HIGH": 16.0, "FLAG": "L",
     "RESULTED_TS": pd.Timestamp("2025-01-02"), "SPECIMEN_TYPE": "BLOOD"},
    {"LAB_ORDER_ID": "LO1", "ORDERED_TS": pd.Timestamp("2025-01-02"), "PRIORITY": "ROUTINE", "TEST_CODE": "WBC",
     "TEST_NAME": "White cells", "VALUE": 7.0, "UNIT": "10^9/L", "REF_LOW": 4.0, "REF_HIGH": 11.0, "FLAG": "N",
     "RESULTED_TS": pd.Timestamp("2025-01-02"), "SPECIMEN_TYPE": "BLOOD"},
])


@pytest.mark.parametrize("role", [r.value for r in Role if r != Role.LAB])
def test_only_lab_staff_can_use_lab_docs(role):
    with pytest.raises(AccessDenied):
        ld.list_documents(Principal("U", "u", role))


def test_upload_rejects_bad_type_before_touching_storage():
    with pytest.raises(ValueError, match="document type"):
        ld.upload(Principal("U", "lab", "LAB_STAFF"), patient_id="P1", doc_type="nope", file_name="a.pdf", data=b"x")


def test_docx_and_csv_reports():
    from docx import Document
    doc = Document(io.BytesIO(ld.report_docx(HDR, RES, "Lab One")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Laboratory Report" in text and "1 outside the reference range" in text
    assert any("Hemoglobin" in c.text for t in doc.tables for r in t.rows for c in r.cells)
    csv = ld.report_csv(RES).decode()
    assert csv.splitlines()[0].startswith("LAB_ORDER_ID") and "Hemoglobin" in csv


def test_patient_id_normalization():
    assert ld.normalize_pid('SYN000123') == 'SYN000123'
    assert ld.normalize_pid(' syn123 ') == 'SYN000123'
    assert ld.normalize_pid('123') == 'SYN000123'
    assert ld.normalize_pid('P000123') == 'P000123'
